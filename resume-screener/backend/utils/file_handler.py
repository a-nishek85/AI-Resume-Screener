import fitz  # PyMuPDF
import pdfplumber
from docx import Document
from pathlib import Path
import re


def extract_text_from_pdf(filepath: str) -> str:
    """Extract text from PDF using PyMuPDF with pdfplumber fallback."""
    text = ""
    try:
        doc = fitz.open(filepath)
        for page in doc:
            text += page.get_text()
        doc.close()
        if len(text.strip()) < 100:
            raise ValueError("Low quality extraction")
        return text.strip()
    except Exception:
        # Fallback to pdfplumber
        try:
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            return text.strip()
        except Exception as e:
            raise ValueError(f"Could not extract text from PDF: {e}")


def extract_text_from_docx(filepath: str) -> str:
    """Extract text from DOCX file."""
    try:
        doc = Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Could not extract text from DOCX: {e}")


def extract_text(filepath: str) -> str:
    """Route to correct extractor based on file extension."""
    ext = Path(filepath).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(filepath)
    elif ext == ".txt":
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def clean_text(text: str) -> str:
    """Clean extracted text."""
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'[^\x00-\x7F]+', ' ', text)
    return text.strip()